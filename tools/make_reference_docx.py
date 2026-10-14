"""Crée assets/note-reference.docx : le modèle de mise en page Word de la note de veille
(styles utilisés par pandoc : titres, corps, tableaux, liens).
Usage : python tools/make_reference_docx.py
"""
import os
import subprocess

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "assets", "note-reference.docx")
BLEU = RGBColor(0x00, 0x00, 0x91)      # bleu institutionnel
ENCRE = RGBColor(0x16, 0x20, 0x2B)
GRIS = RGBColor(0x5A, 0x67, 0x76)
FONT = "Arial"

subprocess.run(["pandoc", "-o", OUT, "--print-default-data-file", "reference.docx"], check=True)
doc = Document(OUT)

for s in doc.sections:
    s.page_height, s.page_width = Cm(29.7), Cm(21.0)
    s.top_margin = s.bottom_margin = Cm(2.0)
    s.left_margin = s.right_margin = Cm(2.2)


STYLES = {s.name: s for s in doc.styles}


def style(name, size, bold=False, color=ENCRE, before=0, after=6, italic=False, align=None, keep=False):
    st = STYLES.get(name)
    if st is None:
        return
    f = st.font
    f.name, f.size, f.bold, f.italic, f.color.rgb = FONT, Pt(size), bold, italic, color
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), FONT)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(a)) is not None:
            del rfonts.attrib[qn(a)]
    if not hasattr(st, "paragraph_format"):
        return
    pf = st.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = 1.15
    if align is not None:
        pf.alignment = align
    if keep:
        pf.keep_with_next = True


style("Normal", 10.5)
style("Body Text", 10.5, after=6, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
style("First Paragraph", 10.5, after=6, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
style("Compact", 10.5, after=2)
style("Title", 20, bold=True, color=BLEU, after=4)
style("Subtitle", 11, color=GRIS, after=12)
style("Heading 1", 16, bold=True, color=BLEU, before=6, after=8, keep=True)
style("Heading 2", 13, bold=True, color=BLEU, before=16, after=6, keep=True)
style("Heading 3", 11.5, bold=True, color=ENCRE, before=10, after=4, keep=True)
style("Heading 4", 10.5, bold=True, color=ENCRE, before=8, after=3, keep=True)
style("Table", 9.5, after=0)
style("Hyperlink", 10.5, color=BLEU)

# Filet sous les titres de niveau 2
h2 = STYLES["Heading 2"].element.get_or_add_pPr()
bdr = OxmlElement("w:pBdr")
bottom = OxmlElement("w:bottom")
for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "2"), ("w:color", "000091")):
    bottom.set(qn(k), v)
bdr.append(bottom)
h2.append(bdr)

# Pied de page : numéro de page
footer = doc.sections[0].footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run()
for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
    if kind:
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
    else:
        el = OxmlElement("w:instrText")
        el.set(qn("xml:space"), "preserve")
        el.text = text
    run._r.append(el)
run.font.size, run.font.color.rgb, run.font.name = Pt(8.5), GRIS, FONT

doc.save(OUT)
print("modèle :", OUT)
